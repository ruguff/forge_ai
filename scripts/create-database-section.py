from pathlib import Path
from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).resolve().parents[1] / 'docs/Section_3_1_Database_Tables.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(20);sec.bottom_margin=Mm(20);sec.left_margin=Mm(20);sec.right_margin=Mm(20)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(12)
    s.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.0
doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
doc.styles['Heading 2'].font.size=Pt(14);doc.styles['Heading 2'].font.bold=True
doc.styles['Caption'].font.bold=True;doc.styles['Caption'].font.italic=False
title=doc.add_paragraph('3.1 Database Tables',style='Title');title.alignment=WD_ALIGN_PARAGRAPH.LEFT

def table(caption,heads,rows,widths):
    p=doc.add_paragraph(caption,style='Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(heads));t.style='Table Grid';t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for c,h,w in zip(t.rows[0].cells,heads,widths):c.width=Mm(w);c.text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,value,w in zip(cells,row,widths):c.width=Mm(w);c.text=value
    for ri,row in enumerate(t.rows):
        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            pr=c._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                n=OxmlElement('w:'+side);n.set(qn('w:w'),'80');n.set(qn('w:type'),'dxa');margins.append(n)
            pr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=(ri==0)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)

heads=['Column','Data type','Key / constraint','Description']
widths=[29,22,44,75]
table('Table 3.1 ProjectForge snapshots',heads,[
 ['project_id','TEXT','Primary key','Unique project identifier supplied by the application.'],
 ['payload','TEXT','NOT NULL','Latest project state, context and build specification serialized as JSON.']],widths)
table('Table 3.2 ProjectForge history',heads,[
 ['id','INTEGER','Primary key; automatically assigned','Identifier of a saved project-state change.'],
 ['project_id','TEXT','Nullable; no foreign-key constraint','Identifier of the associated project.'],
 ['payload','TEXT','NOT NULL','Project snapshot serialized as JSON at the time of saving.']],widths)
table('Table 3.3 ProjectForge events',heads,[
 ['id','INTEGER','Primary key; AUTOINCREMENT','Global event sequence number used for ordered replay.'],
 ['project_id','TEXT','NOT NULL; no foreign-key constraint','Identifier of the project associated with the event.'],
 ['payload','TEXT','NOT NULL','Execution-event details serialized as JSON.']],widths)

doc.add_page_break()
doc.add_paragraph('3.1 Database Tables Continued',style='Title')
table('Table 3.4 Generated NoteKeeper notes',heads,[
 ['id','INTEGER','Primary key; NOT NULL; automatically assigned','Unique note identifier.'],
 ['title','TEXT','NOT NULL; default empty string','Note title.'],
 ['body','TEXT','NOT NULL; default empty string','Note content.']],widths)
table('Table 3.5 ProjectForge indexes',['Index','Table','Columns','Purpose'],[
 ['sqlite_autoindex_\nsnapshots_1','snapshots','project_id','Automatically enforces primary-key uniqueness.'],
 ['events_project','events','project_id, id','Supports retrieval of a project’s events in sequence.']],[43,25,32,70])
table('Table 3.6 Logical project relationships',['Parent','Related column','Cardinality','Enforcement'],[
 ['snapshots.\nproject_id','history.project_id','One to many','Application-maintained relationship; no declared foreign key.'],
 ['snapshots.\nproject_id','events.project_id','One to many','Application-maintained relationship; no declared foreign key.']],[35,40,27,68])
table('Table 3.7 Database scope',['Item','Implemented scope'],[
 ['ProjectForge storage','SQLite file: backend/data/projects.sqlite3; application tables: snapshots, history and events.'],
 ['Generated application storage','Separate SQLite file: generated NoteKeeper backend/notes.db; application table: notes.'],
 ['Internal SQLite table','sqlite_sequence maintains AUTOINCREMENT counters; it is not an application entity.'],
 ['Constraint interpretation','NOT NULL rejects NULL but does not reject empty strings. Logical project relationships are not database-enforced foreign keys.']],[48,122])
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='ProjectForge Section 3.1 Database Tables'
doc.core_properties.subject='Implemented SQLite data dictionary'
for style in doc.styles:
    fonts=style.element.find(qn('w:rPr'))
    if fonts is not None:
        rf=fonts.find(qn('w:rFonts'))
        if rf is not None:
            for attr in list(rf.attrib):
                if 'Theme' in attr or 'theme' in attr:del rf.attrib[attr]
            for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for p in doc.paragraphs:
    for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for run in p.runs:
        run.font.name='Times New Roman'
        if p.style.name=='Title':run.font.size=Pt(16);run.bold=True
doc.save(out)
print(out)
