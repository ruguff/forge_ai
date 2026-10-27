from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / 'ProjectForge_Sections_4_to_6_1.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Inches(0.85)
sec.left_margin = Inches(1)
sec.right_margin = Inches(0.85)
sec.footer_distance = Inches(0.4)

def set_font(style, size=12, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + a), 'Times New Roman')
    for a in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + a), None)

for s in doc.styles:
    if s.type in (1, 2):
        set_font(s)
normal = doc.styles['Normal']
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 16), ('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 14)]:
    s = doc.styles[name]
    set_font(s, size, True)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    s.paragraph_format.space_before = Pt(10)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
for name in ['Caption', 'List Bullet', 'List Number']:
    set_font(doc.styles[name])
    doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles['Caption'].font.italic = True
doc.styles['Caption'].paragraph_format.space_after = Pt(9)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run()
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
r._r.addnext(fld)
doc.core_properties.title = 'ProjectForge AI Sections 4 to 6.1'
doc.core_properties.subject = 'Implementation, Conclusion and References'
doc.core_properties.author = 'Panchal Het Vipul Kumar'
doc.core_properties.keywords = 'ProjectForge AI, final year project, implementation'

def inline(p, text):
    for i, token in enumerate(re.split(r'(\*\*.*?\*\*)', text)):
        if not token: continue
        run = p.add_run(token[2:-2] if token.startswith('**') else token)
        if token.startswith('**'): run.bold = True
    return p

def para(text, style=None):
    p = doc.add_paragraph(style=style)
    inline(p, text)
    return p

def table(rows):
    n = len(rows[0])
    t = doc.add_table(rows=1, cols=n)
    t.autofit = False
    widths = [1.8, 1.65, 2.97] if n == 3 else [2.25, 4.17]
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        v = OxmlElement('w:' + k)
        v.set(qn('w:val'),'single');v.set(qn('w:sz'),'6');v.set(qn('w:color'),'D9D9D9')
        borders.append(v)
    pr.append(borders)
    mar = OxmlElement('w:tblCellMar')
    for k, val in [('top','85'),('bottom','85'),('left','100'),('right','100')]:
        v=OxmlElement('w:'+k);v.set(qn('w:w'),val);v.set(qn('w:type'),'dxa');mar.append(v)
    pr.append(mar)
    for idx, vals in enumerate(rows):
        row = t.rows[0] if idx == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx == 0: trpr.append(OxmlElement('w:tblHeader'))
        for c,w,text in zip(row.cells,widths,vals):
            c.width = Inches(w);c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            run=p.add_run(text);run.bold=idx==0
            if idx==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');c._tc.get_or_add_tcPr().append(sh)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=(ROOT/'sections_4_to_6_1.md').read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i]
    if not line.strip(): i+=1;continue
    if line.startswith('```'):
        if doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
        code=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            code.append(lines[i]);i+=1
        for j, text in enumerate(code):
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1
            p.paragraph_format.keep_with_next=j<len(code)-1
            p.paragraph_format.keep_together=True
            indent=len(text)-len(text.lstrip(' '))
            p.paragraph_format.left_indent=Inches(0.15+indent*0.045)
            p.add_run(text.lstrip(' '))
        p.paragraph_format.space_after=Pt(8)
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[c.strip() for c in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',c) for c in cells): rows.append(cells)
            i+=1
        table(rows);continue
    elif line.startswith('# '):
        p=para(line[2:], 'Heading 1')
        if line.startswith('# 5.') or line.startswith('# 6.'):p.paragraph_format.page_break_before=True
    elif line.startswith('## '):
        p=para(line[3:], 'Heading 2')
    elif line.startswith('### '): para(line[4:], 'Heading 3')
    elif line.startswith('[[SCREENSHOT:'):
        p=para('['+line[len('[[SCREENSHOT: '):-2]+']')
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.italic=True
    elif line.startswith('Figure '):para(line,'Caption')
    elif line.startswith('Table '):
        p=para(line,'Caption');p.paragraph_format.keep_with_next=True
    elif line.startswith('- '):
        p=para(line[2:],'List Bullet');p.paragraph_format.space_after=Pt(3)
    elif re.match(r'^\d+\. ',line):
        p=para(line)
        p.paragraph_format.left_indent=Inches(0.24)
        p.paragraph_format.first_line_indent=Inches(-0.24)
        p.paragraph_format.space_after=Pt(5)
    else:
        if line.startswith('**Source:**') and doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p=para(line)
        if line.startswith('**') and line.endswith('**'):p.paragraph_format.keep_with_next=True
    i+=1

# Explicitly remove theme font inheritance from all runs and preserve requested sizes.
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]+list(sec.footer.paragraphs):
    for r in p.runs:
        r.font.name='Times New Roman'
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for k in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+k),'Times New Roman')
        for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+k),None)
doc.save(OUT)
print(OUT)
print(f'Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}')
